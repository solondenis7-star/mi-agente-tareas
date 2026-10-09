
import os
import re
import tempfile
import gradio as gr
from docx import Document
from docx.shared import Pt
from google import genai

def resolver_word(archivo):
    if not archivo:
        raise gr.Error("Primero selecciona un archivo Word.")

    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        raise gr.Error("Falta configurar la clave de Gemini.")

    try:
        documento = Document(archivo)
        partes = []

        for p in documento.paragraphs:
            if p.text.strip():
                partes.append(p.text)

        for tabla in documento.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(c.text for c in fila.cells))

        texto = "\n".join(partes)
        if not texto.strip():
            raise gr.Error("El documento no contiene texto legible.")

        cliente = genai.Client(api_key=api_key)
        prompt = f"""
Resuelve esta ficha escolar de computación para un estudiante
de secundaria.

Reglas:
- Responde todas las preguntas en el orden original.
- Usa lenguaje claro y apropiado para secundaria.
- No inventes nombres, integrantes ni actividades realizadas.
- Si falta un dato personal, deja el espacio en blanco.
- Si una actividad requiere información que no aparece,
  indica qué debe completar el estudiante.
- Conserva los títulos y la numeración.
- Responde todas las secciones.
- No uses marcas Markdown como ** o ##.
- Entrega las preguntas y sus respuestas.

DOCUMENTO ORIGINAL:
{texto}
"""
        respuesta = cliente.models.generate_content(
            model="gemini-2.5-flash",
            contents=prompt
        )
        contenido = respuesta.text or ""
        if not contenido.strip():
            raise gr.Error("La IA no devolvió respuestas. Inténtalo otra vez.")

        salida = Document()
        estilo = salida.styles["Normal"]
        estilo.font.name = "Arial"
        estilo.font.size = Pt(11)

        for linea in contenido.splitlines():
            p = salida.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(linea)
            run.font.name = "Arial"
            run.font.size = Pt(11)

        archivo_salida = tempfile.NamedTemporaryFile(
            suffix=".docx", delete=False
        )
        archivo_salida.close()
        salida.save(archivo_salida.name)
        return archivo_salida.name

    except gr.Error:
        raise
    except Exception as e:
        raise gr.Error(
            "Ocurrió un error al resolver el archivo. "
            "Revisa que sea un .docx y vuelve a intentarlo."
        )

with gr.Blocks(title="Mi Agente de Tareas") as app:
    gr.Markdown("# 📚 Mi Agente de Tareas")
    gr.Markdown(
        "Sube una ficha de Word, resuelve sus preguntas "
        "y descarga las respuestas en Arial 11."
    )
    archivo = gr.File(
        label="1. Selecciona la ficha del profesor",
        file_types=[".docx"],
        type="filepath"
    )
    boton = gr.Button("🤖 Resolver tarea", variant="primary")
    resultado = gr.File(label="2. Descarga tu tarea resuelta")
    boton.click(resolver_word, inputs=archivo, outputs=resultado)

app.launch(server_name="0.0.0.0", server_port=int(os.environ.get("PORT", 7860)))
